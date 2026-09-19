import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "cleanup_sample.xlsx"

HEADERS = [
    "Record_ID",
    "Customer",
    "Email",
    "Created",
    "Status",
    "Amount",
    "Currency",
    "Owner",
    "Notes",
]

RAW_ROWS = [
    ["A-1001", "  Alice Chen ", "alice@example.com", "2026-09-01", "paid", 120.0, "USD", "Maya", "ok"],
    ["A-1002", "Bob Li", "BOB@EXAMPLE.COM", "2026-09-01", "pending", 80.0, "USD", "Leo", ""],
    ["A-1002", "Bob Li", "BOB@EXAMPLE.COM", "2026-09-01", "pending", 80.0, "USD", "Leo", "duplicate upload"],
    ["A-1003", "Carla Gomez", "", "2026-09-02", "paid", 240.0, "USD", "Maya", "missing email"],
    ["A-1004", "David Kim", "david@example.com", "2026-09-02", "refunded", -25.0, "USD", "Leo", "negative amount"],
    ["A-1005", "Elena Petrova", "elena@example.com", "09/03/2026", "PAID", 310.0, "USD", "Nora", "date needs normalization"],
    ["A-1006", "Fatima Noor", "fatima@example.com", "2026-09-03", "paid", 175.0, "USD", "Maya", ""],
    ["A-1007", "George Tan", "george@example.com", "not a date", "pending", 60.0, "USD", "Leo", "invalid date"],
    ["A-1008", "Hana Ito", "hana@example.com", "2026-09-04", "unknown", 95.0, "USD", "Nora", "unknown status"],
    ["A-1009", "Ivan Petrov", "ivan@example.com", "2026-09-04", "paid", 420.0, "USD", "Maya", ""],
    ["A-1010", "Joy Okafor", "joy@example.com", "2026-09-05", "pending", 135.0, "USD", "Leo", ""],
    ["A-1011", "Kai Muller", "kai@example.com", "2026-09-05", "paid", 210.0, "USD", "Nora", ""],
    ["A-1012", "Lina Haddad", "LINA@EXAMPLE.COM", "2026-09-05", "paid", 150.0, "USD", "Maya", ""],
    ["A-1013", "Mina Park", "mina@example.com", "2026-09-06", "pending", 0.0, "USD", "Nora", "zero amount"],
    ["A-1014", "Noah Brown", "noah@example.com", "2026-09-06", "paid", 280.0, "USD", "Leo", ""],
]

LOOKUP_ROWS = [
    ["Maya", "maya@internal.example"],
    ["Leo", "leo@internal.example"],
    ["Nora", "nora@internal.example"],
]


def normalize_date(value):
    if not isinstance(value, str):
        return value
    if len(value) == 10 and value[4] == "-" and value[7] == "-":
        return value
    if len(value) == 10 and value[2] == "/" and value[5] == "/":
        month, day, year = value.split("/")
        return f"{year}-{month}-{day}"
    return value


def classify(record):
    record_id, customer, email, created, status, amount, _, owner, _ = record
    reasons = []
    if not email:
        reasons.append("Missing email")
    if amount <= 0:
        reasons.append("Amount is not positive")
    if normalize_date(created) == created and not re.fullmatch(
        r"\d{4}-\d{2}-\d{2}",
        str(created),
    ):
        reasons.append("Invalid date")
    if status.lower() not in {"paid", "pending", "refunded"}:
        reasons.append("Unknown status")
    return record_id, customer, email, created, status, amount, owner, reasons


def style_header(worksheet, start_col, end_col, row=1):
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet.iter_rows(
        min_row=row,
        max_row=row,
        min_col=start_col,
        max_col=end_col,
    ):
        for item in cell:
            item.fill = fill
            item.font = Font(color="FFFFFF", bold=True)
            item.alignment = Alignment(horizontal="center")


def build_workbook():
    workbook = Workbook()
    raw = workbook.active
    raw.title = "Raw_Data"
    raw.append(HEADERS)
    for row in RAW_ROWS:
        raw.append(row)
    raw.freeze_panes = "A2"
    raw.auto_filter.ref = f"A1:I{raw.max_row}"
    raw.conditional_formatting.add(
        f"F2:F{raw.max_row}",
        CellIsRule(
            operator="lessThanOrEqual",
            formula=["0"],
            fill=PatternFill("solid", fgColor="FCE4D6"),
            font=Font(color="9C0006"),
        ),
    )

    cleaned = workbook.create_sheet("Clean_Records")
    clean_headers = HEADERS + ["Owner_Email", "Quality_Check"]
    cleaned.append(clean_headers)
    duplicate_ids = set()
    seen_ids = set()
    exceptions = []
    clean_count = 0

    for row_number, record in enumerate(RAW_ROWS, start=2):
        record_id, customer, email, created, status, amount, owner, reasons = (
            classify(record)
        )
        if record_id in seen_ids:
            duplicate_ids.add(record_id)
            exceptions.append(
                ["Raw_Data", row_number, record_id, "Duplicate Record_ID", "Do not load until reviewed"]
            )
            continue
        seen_ids.add(record_id)
        if reasons:
            for reason in reasons:
                exceptions.append(
                    ["Raw_Data", row_number, record_id, reason, "Review before loading"]
                )
            continue

        target_row = cleaned.max_row + 1
        cleaned.append(
            [
                record_id,
                customer.strip(),
                email.strip().lower(),
                normalize_date(created),
                status.lower(),
                amount,
                record[6],
                owner,
                record[8].strip(),
                f'=IFERROR(VLOOKUP(H{target_row},Lookup!$A:$B,2,FALSE),"REVIEW")',
                '=IF(COUNTA(A{0}:I{0})=9,"PASS","REVIEW")'.format(target_row),
            ]
        )
        clean_count += 1

    cleaned.freeze_panes = "A2"
    cleaned.auto_filter.ref = f"A1:K{cleaned.max_row}"
    status_validation = DataValidation(
        type="list",
        formula1='"paid,pending,refunded"',
        allow_blank=False,
    )
    cleaned.add_data_validation(status_validation)
    status_validation.add(f"E2:E{cleaned.max_row}")
    cleaned.conditional_formatting.add(
        f"K2:K{cleaned.max_row}",
        FormulaRule(
            formula=[f'$K2<>"PASS"'],
            fill=PatternFill("solid", fgColor="FFC7CE"),
            font=Font(color="9C0006"),
        ),
    )
    for row in cleaned.iter_rows(min_row=2, min_col=6, max_col=6):
        for cell in row:
            cell.number_format = '$#,##0.00'

    exception_sheet = workbook.create_sheet("Exceptions")
    exception_sheet.append(
        ["Source_Sheet", "Source_Row", "Record_ID", "Issue", "Required_Action"]
    )
    for row in exceptions:
        exception_sheet.append(row)
    exception_sheet.freeze_panes = "A2"
    exception_sheet.auto_filter.ref = f"A1:E{exception_sheet.max_row}"

    lookup = workbook.create_sheet("Lookup")
    lookup.append(["Owner", "Owner_Email"])
    for row in LOOKUP_ROWS:
        lookup.append(row)
    lookup.freeze_panes = "A2"

    summary = workbook.create_sheet("Weekly_Summary")
    summary.append(["Owner", "Paid_Total", "Pending_Total", "Record_Count", "Review_Count"])
    for row_number, owner in enumerate(["Maya", "Leo", "Nora"], start=2):
        summary.append(
            [
                owner,
                f'=SUMIFS(Clean_Records!$F:$F,Clean_Records!$H:$H,$A{row_number},Clean_Records!$E:$E,"paid")',
                f'=SUMIFS(Clean_Records!$F:$F,Clean_Records!$H:$H,$A{row_number},Clean_Records!$E:$E,"pending")',
                f'=COUNTIF(Clean_Records!$H:$H,$A{row_number})',
                f'=COUNTIFS(Clean_Records!$H:$H,$A{row_number},Clean_Records!$K:$K,"REVIEW")',
            ]
        )
    summary.append(
        [
            "TOTAL",
            "=SUM(B2:B4)",
            "=SUM(C2:C4)",
            "=SUM(D2:D4)",
            "=SUM(E2:E4)",
        ]
    )
    for row in summary.iter_rows(min_row=2, min_col=2, max_col=3):
        for cell in row:
            cell.number_format = '$#,##0.00'
    summary.freeze_panes = "A2"
    style_header(summary, 1, 5)

    chart = BarChart()
    chart.title = "Paid Total by Owner"
    chart.y_axis.title = "USD"
    chart.x_axis.title = "Owner"
    data = Reference(summary, min_col=2, min_row=1, max_row=4)
    categories = Reference(summary, min_col=1, min_row=2, max_row=4)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.height = 7
    chart.width = 12
    summary.add_chart(chart, "G2")

    rules = workbook.create_sheet("Rules")
    rules.append(["Area", "Rule", "Result"])
    rules.append(["Identity", "Record_ID must be unique", "Duplicates are quarantined"])
    rules.append(["Customer", "Trim outer whitespace", "Normalized customer name"])
    rules.append(["Email", "Require non-empty, lower-case address", "Missing values are quarantined"])
    rules.append(["Date", "Use YYYY-MM-DD", "Convertible US dates are normalized"])
    rules.append(["Status", "paid, pending, refunded", "Unknown values are quarantined"])
    rules.append(["Amount", "Must be numeric and positive", "Zero and negative values are quarantined"])
    rules.append(["Owner", "Must exist in Lookup", "VLOOKUP returns REVIEW when unmatched"])
    rules.append(["Audit", "Every change or exception is traceable", "Exceptions sheet records source row"])
    rules.freeze_panes = "A2"
    style_header(rules, 1, 3)

    for worksheet in workbook.worksheets:
        for column_cells in worksheet.columns:
            width = max(len(str(cell.value or "")) for cell in column_cells) + 3
            worksheet.column_dimensions[column_cells[0].column_letter].width = min(
                max(width, 12),
                42,
            )

    style_header(raw, 1, len(HEADERS))
    style_header(cleaned, 1, len(clean_headers))
    style_header(exception_sheet, 1, 5)
    style_header(lookup, 1, 2)
    workbook.save(OUTPUT)

    return {
        "clean_records": clean_count,
        "exceptions": len(exceptions),
        "duplicates": len(duplicate_ids),
    }


def verify_workbook(summary):
    workbook = load_workbook(OUTPUT, data_only=False)
    assert workbook.sheetnames == [
        "Raw_Data",
        "Clean_Records",
        "Exceptions",
        "Lookup",
        "Weekly_Summary",
        "Rules",
    ]
    assert summary == {"clean_records": 9, "exceptions": 6, "duplicates": 1}
    assert workbook["Clean_Records"]["J2"].value.startswith("=IFERROR(VLOOKUP")
    assert workbook["Clean_Records"]["K2"].value.startswith("=IF(COUNTA")
    assert workbook["Weekly_Summary"]["B5"].value == "=SUM(B2:B4)"
    assert len(workbook["Clean_Records"].data_validations.dataValidation) == 1
    assert len(workbook["Raw_Data"].conditional_formatting) == 1
    print(
        f"verified {OUTPUT.name}: {summary['clean_records']} clean, "
        f"{summary['exceptions']} exceptions, {summary['duplicates']} duplicate"
    )


if __name__ == "__main__":
    verify_workbook(build_workbook())
